import csv
import hashlib
import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime

import openpyxl

from .errors import ImportFileError

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLS = 60
HEADER_SCAN_ROWS = 10
ALLOWED_EXTENSIONS = (".xlsx", ".csv")


@dataclass
class ParsedFile:
    sheets: list
    sheet: str
    header_row: int
    headers: list
    rows: list          # [{"n": sheet row number, "v": [values]}]
    sha256: str


def _cell(v):
    """Make a cell JSON-safe and tidy: dates -> ISO text, 12.0 -> 12, blank -> None."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, (int, float)):
        return v
    text = str(v).strip()
    return text or None


def _is_blank(cells):
    return all(c is None for c in cells)


def _find_header(numbered_rows, forced):
    """Return the index (into numbered_rows) of the header row."""
    if forced:
        for i, (n, _) in enumerate(numbered_rows):
            if n == forced:
                return i
        raise ImportFileError(f"Row {forced} was not found in the sheet.")
    for i, (n, cells) in enumerate(numbered_rows[:HEADER_SCAN_ROWS]):
        filled = [c for c in cells if c is not None]
        texts = [c for c in filled if isinstance(c, str)]
        if len(filled) >= 3 and len(texts) * 2 >= len(filled):
            return i
    raise ImportFileError("Could not find a header row (a row with at least 3 column titles) in the first 10 rows.")


def _clean_headers(cells):
    last = max((i for i, c in enumerate(cells) if c is not None), default=-1)
    headers, seen = [], {}
    for i, c in enumerate(cells[: last + 1]):
        name = str(c).strip() if c is not None else f"Column {i + 1}"
        key = name.casefold()
        seen[key] = seen.get(key, 0) + 1
        headers.append(name if seen[key] == 1 else f"{name} ({seen[key]})")
    return headers


def _build(numbered_rows, forced_header_row):
    h = _find_header(numbered_rows, forced_header_row)
    headers = _clean_headers(numbered_rows[h][1])
    if not headers:
        raise ImportFileError("The header row is empty.")
    if len(headers) > MAX_COLS:
        raise ImportFileError(f"Too many columns ({len(headers)}). The limit is {MAX_COLS}.")
    rows = []
    for n, cells in numbered_rows[h + 1:]:
        values = (list(cells) + [None] * len(headers))[: len(headers)]
        if _is_blank(values):
            continue
        rows.append({"n": n, "v": values})
        if len(rows) > MAX_ROWS:
            raise ImportFileError(f"Too many rows. The limit is {MAX_ROWS} rows per import; split the file.")
    if not rows:
        raise ImportFileError("The sheet has a header row but no data rows.")
    return numbered_rows[h][0], headers, rows


def _read_xlsx(content, sheet):
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError, openpyxl.utils.exceptions.InvalidFileException):
        raise ImportFileError("This does not look like a valid .xlsx file. (Old .xls files are not supported; save as .xlsx.)")
    sheets = list(wb.sheetnames)
    if not sheets:
        raise ImportFileError("The workbook has no sheets.")
    chosen = sheet or sheets[0]
    if chosen not in sheets:
        raise ImportFileError(f"Sheet '{chosen}' not found. Available sheets: {', '.join(sheets)}.")
    numbered = []
    for n, row in enumerate(wb[chosen].iter_rows(values_only=True), start=1):
        numbered.append((n, [_cell(c) for c in row]))
        if n > MAX_ROWS + 200:                     # stop reading absurdly large sheets early
            raise ImportFileError(f"Too many rows. The limit is {MAX_ROWS} rows per import; split the file.")
    wb.close()
    return sheets, chosen, numbered


def _read_csv(content):
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportFileError("Could not read the file's text encoding. Save it as UTF-8 CSV.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    numbered = []
    for n, row in enumerate(csv.reader(io.StringIO(text), dialect), start=1):
        numbered.append((n, [_cell(c) for c in row]))
        if n > MAX_ROWS + 200:
            raise ImportFileError(f"Too many rows. The limit is {MAX_ROWS} rows per import; split the file.")
    return ["CSV"], "CSV", numbered


def parse_upload(filename, content, sheet=None, header_row=None):
    name = (filename or "").lower()
    if not name.endswith(ALLOWED_EXTENSIONS):
        raise ImportFileError("Only .xlsx and .csv files can be imported.")
    if len(content) > MAX_BYTES:
        raise ImportFileError(f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB.")
    if not content:
        raise ImportFileError("The file is empty.")
    sheets, chosen, numbered = _read_xlsx(content, sheet) if name.endswith(".xlsx") else _read_csv(content)
    header_n, headers, rows = _build(numbered, header_row)
    return ParsedFile(sheets, chosen, header_n, headers, rows, hashlib.sha256(content).hexdigest())
